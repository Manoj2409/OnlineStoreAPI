import csv
import json


def read_excel_data(filepath, sheetname):
    from openpyxl import load_workbook

    workbook = load_workbook(filepath)
    sheet = workbook[sheetname]
    data = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        data.append(row)

    return data


def read_json_data(filepath):
    with open(filepath, "r") as file:
        data_list = json.load(file)

    return [(item,) for item in data_list]


def read_csv_data(filepath):
    data = []

    with open(filepath, "r") as file:
        reader = csv.reader(file)
        next(reader)

        for row in reader:
            data.append(tuple(row))

    return data
